from __future__ import annotations

import json
from pathlib import Path

from docx import Document


FILES = [
    Path(r"C:\Users\hq362\Downloads\Mini-OGAS_demo_frontend_report.docx"),
    Path(r"C:\Users\hq362\Downloads\Mini-OGAS_real_management_frontend_report.docx"),
]
OUT = Path(r"D:\New project\mini-ogas\.runtime\frontend-docs")


def extract_doc(path: Path) -> dict:
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)
    return {"file": str(path), "paragraphs": paragraphs, "tables": tables}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    docs = [extract_doc(path) for path in FILES]
    (OUT / "frontend_docs.json").write_text(json.dumps(docs, ensure_ascii=False, indent=2), encoding="utf-8")
    for item in docs:
        txt = ["# " + Path(item["file"]).name, ""]
        txt.extend(item["paragraphs"])
        for index, table in enumerate(item["tables"], start=1):
            txt.append("")
            txt.append(f"## Table {index}")
            for row in table:
                txt.append(" | ".join(row))
        (OUT / (Path(item["file"]).stem + ".md")).write_text("\n".join(txt), encoding="utf-8")
    print(OUT / "frontend_docs.json")


if __name__ == "__main__":
    main()
