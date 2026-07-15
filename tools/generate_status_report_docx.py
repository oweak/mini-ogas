from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "mini-ogas-current-status-report-2026-05-22.md"
TARGET = ROOT / "docs" / "Mini-OGAS_分布式系统现状与架构报告_2026-05-23.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text.strip())
    run.bold = bold
    run.font.size = Pt(9)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    for name, size, color in [
        ("Title", 22, "17365D"),
        ("Heading 1", 16, "1D4ED8"),
        ("Heading 2", 13, "17406D"),
        ("Heading 3", 11.5, "334155"),
    ]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)


def format_chinese_date(date_text: str) -> str:
    match = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", date_text)
    if match:
        year, month, day = match.groups()
        return f"{year}年{int(month)}月{int(day)}日"
    return date_text


def get_source_meta(text: str, key: str) -> str | None:
    pattern = rf"^{re.escape(key)}：\s*(.+)$"
    for line in text.splitlines():
        match = re.match(pattern, line)
        if match:
            return match.group(1).strip()
    return None


def set_run_style(run, font_name: str = "Microsoft YaHei", size: Pt | None = None, color: RGBColor | None = None, bold: bool = False) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = size
    if color is not None:
        run.font.color.rgb = color
    run.bold = bold


def add_cover(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(80)
    run = title.add_run("Mini-OGAS 分布式系统\n现状与架构报告")
    set_run_style(run, size=Pt(24), color=RGBColor(23, 54, 93), bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(20)
    run = subtitle.add_run("父节点 / 云端子节点 / DeepSeek API / 自然语言命令网关")
    set_run_style(run, size=Pt(12), color=RGBColor(71, 85, 105))

    markdown = SOURCE.read_text(encoding="utf-8")
    report_date = get_source_meta(markdown, "生成日期") or "2026-05-23"
    report_date = format_chinese_date(report_date)

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.style = "Table Grid"
    rows = [
        ("生成日期", report_date),
        ("项目目录", r"D:\New project\mini-ogas"),
        ("项目阶段", "复试展示型 MVP / 本机 + 云端混合部署"),
        ("报告内容", "操作记录、当前现状、系统架构、接口与后续计划"),
    ]
    for row, (key, value) in zip(meta.rows, rows):
        set_cell_text(row.cells[0], key, True)
        set_cell_shading(row.cells[0], "EAF2F8")
        set_cell_text(row.cells[1], value)

    doc.add_page_break()


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EFF6FF")
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_style(run)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph()


def add_code_block(doc: Document, code: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.right_indent = Cm(0.15)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(code.rstrip())
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)


def is_table(lines: list[str], index: int) -> bool:
    return (
        index + 1 < len(lines)
        and lines[index].strip().startswith("|")
        and re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", lines[index + 1])
    )


def parse_table_block(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        line = lines[index].strip().strip("|")
        cells = [cell.strip().replace("`", "") for cell in line.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            text = row[c_idx] if c_idx < len(row) else ""
            set_cell_text(table.cell(r_idx, c_idx), text, bold=(r_idx == 0))
            if r_idx == 0:
                set_cell_shading(table.cell(r_idx, c_idx), "D9EAF7")
    doc.add_paragraph()


def add_markdown_content(doc: Document, text: str) -> None:
    lines = text.splitlines()
    in_code = False
    code_lines: list[str] = []
    skip_title = True
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if is_table(lines, i):
            rows, i = parse_table_block(lines, i)
            add_table(doc, rows)
            continue

        if stripped.startswith("# "):
            if skip_title:
                skip_title = False
            else:
                doc.add_heading(stripped[2:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue

        if stripped.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            run = paragraph.add_run(clean_inline(stripped[2:]))
            set_run_style(run)
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            paragraph = doc.add_paragraph(style="List Number")
            run = paragraph.add_run(clean_inline(re.sub(r"^\d+\.\s+", "", stripped)))
            set_run_style(run)
            i += 1
            continue

        if stripped.startswith(">"):
            add_callout(doc, clean_inline(stripped.lstrip("> ").strip()))
            i += 1
            continue

        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.18
        run = paragraph.add_run(clean_inline(stripped))
        set_run_style(run)
        i += 1


def clean_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    return text


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Mini-OGAS 分布式系统现状与架构报告")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 116, 139)


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    style_document(doc)
    add_cover(doc)
    add_markdown_content(doc, markdown)
    add_footer(doc)
    doc.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
