"""Convert the Mini-OGAS technical report from Markdown to .docx."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor

doc = Document()

# -- Page setup --
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# -- Style tweaks --
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

for level in range(1, 4):
    heading_style = doc.styles[f'Heading {level}']
    heading_style.font.name = 'Microsoft YaHei'
    heading_style.font.color.rgb = RGBColor(0x1f, 0x29, 0x33)

with open('docs/Mini-OGAS_技术报告_2026-05-30.md', 'r', encoding='utf-8') as f:
    lines = f.readlines()

i = 0
in_code_block = False
code_lines: list[str] = []
in_table = False
table_rows: list[list[str]] = []

while i < len(lines):
    line = lines[i]

    # Code blocks
    if line.startswith('```'):
        if in_code_block:
            p = doc.add_paragraph()
            p.style = doc.styles['Normal']
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            pf = p.paragraph_format
            pf.left_indent = Cm(0.5)
            pf.space_before = Pt(4)
            pf.space_after = Pt(4)
            code_lines = []
            in_code_block = False
        else:
            in_code_block = True
        i += 1
        continue

    if in_code_block:
        code_lines.append(line.rstrip('\n'))
        i += 1
        continue

    # Tables
    if line.startswith('|') and line.strip().endswith('|'):
        if not in_table:
            in_table = True
            table_rows = []
        # skip separator rows like |---|---|
        if not re.match(r'^\|[\s\-:|]+\|$', line.strip()):
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            table_rows.append(cells)
        i += 1
        continue
    elif in_table:
        # Table ended — flush it
        if table_rows:
            table = doc.add_table(rows=len(table_rows), cols=max(len(r) for r in table_rows))
            table.style = 'Light Grid Accent 1'
            for ri, row in enumerate(table_rows):
                for ci, cell_text in enumerate(row):
                    if ci < len(table.rows[ri].cells):
                        cell = table.rows[ri].cells[ci]
                        cell.text = cell_text
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.font.size = Pt(9)
            doc.add_paragraph()  # spacer
        in_table = False
        table_rows = []
        # don't increment i — reprocess this line
        continue

    # Headings
    if line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=3)
    elif line.startswith('# '):
        doc.add_heading(line[2:].strip(), level=1)
    # Horizontal rule
    elif line.strip() == '---':
        doc.add_paragraph('_' * 60)
    # Blockquote
    elif line.startswith('> '):
        p = doc.add_paragraph(line[2:].strip())
        p.paragraph_format.left_indent = Cm(1)
        for run in p.runs:
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x4a, 0x55, 0x68)
    # Unordered list
    elif line.startswith('- ') or line.startswith('  - '):
        text = line.strip().lstrip('- ').strip()
        # Inline code
        text = re.sub(r'`([^`]+)`', r'\1', text)
        # Bold
        text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
        # Remove HTML tags
        text = re.sub(r'<[^>]+>', '', text)
        p = doc.add_paragraph(text, style='List Bullet')
    # Ordered list
    elif re.match(r'^\d+\.\s', line):
        text = re.sub(r'^\d+\.\s', '', line.strip())
        text = re.sub(r'`([^`]+)`', r'\1', text)
        text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
        text = re.sub(r'<[^>]+>', '', text)
        p = doc.add_paragraph(text, style='List Number')
    # Empty line
    elif line.strip() == '':
        pass  # skip
    # Regular paragraph
    else:
        text = line.strip()
        # Skip image links, HTML comments
        if text.startswith('![') or text.startswith('<!--'):
            i += 1
            continue
        # Remove markdown formatting
        text = re.sub(r'`([^`]+)`', r'\1', text)
        text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
        text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
        text = re.sub(r'<[^>]+>', '', text)
        if text:
            p = doc.add_paragraph(text)
            p.paragraph_format.space_after = Pt(6)

    i += 1

# Flush any remaining table
if in_table and table_rows:
    table = doc.add_table(rows=len(table_rows), cols=max(len(r) for r in table_rows))
    table.style = 'Light Grid Accent 1'
    for ri, row in enumerate(table_rows):
        for ci, cell_text in enumerate(row):
            if ci < len(table.rows[ri].cells):
                table.rows[ri].cells[ci].text = cell_text

output_path = 'docs/Mini-OGAS_技术报告_2026-05-30.docx'
doc.save(output_path)
print(f'Saved: {output_path}')
